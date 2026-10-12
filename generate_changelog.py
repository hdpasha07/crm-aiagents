import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets the background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets padding for a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    Sets specific borders for a cell.
    kwargs: top, bottom, left, right with dict(val='single', sz='12', color='1B365D')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', 'auto'))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def add_callout_box(doc, text, title="KEY INSIGHT", border_color="1B365D", fill_color="F0F4F8"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)

    cell = tbl.cell(0, 0)
    set_cell_background(cell, fill_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=160)
    set_cell_border(cell, left={'val': 'single', 'sz': 24, 'color': border_color},
                          top={'val': 'nil'}, right={'val': 'nil'}, bottom={'val': 'nil'})

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title}: ")
    run_t.bold = True
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(10.5)
    run_t.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    run_body = p.add_run(text)
    run_body.font.name = "Calibri"
    run_body.font.size = Pt(10)
    run_body.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    # Add trailing space after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)

    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    set_cell_border(cell, 
                    left={'val': 'single', 'sz': 6, 'color': 'D0D5DD'},
                    right={'val': 'single', 'sz': 6, 'color': 'D0D5DD'},
                    top={'val': 'single', 'sz': 6, 'color': 'D0D5DD'},
                    bottom={'val': 'single', 'sz': 6, 'color': 'D0D5DD'})

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x1E, 0x1E, 0x1E)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def build_changelog():
    doc = Document()

    # Configure Margins
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.9)
        sec.right_margin = Inches(0.9)

    # Document Header / Banner Table
    header_table = doc.add_table(rows=1, cols=1)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    header_table.columns[0].width = Inches(6.7)

    h_cell = header_table.cell(0, 0)
    set_cell_background(h_cell, "1B365D")  # Deep Navy
    set_cell_margins(h_cell, top=240, bottom=240, left=240, right=240)

    title_p = h_cell.paragraphs[0]
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("Customer Support Agent — Technical Changelog")
    run_title.bold = True
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    sub_p = h_cell.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(0)
    run_sub = sub_p.add_run("LangGraph State Persistence Diagnostic & Resolution Report")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = RGBColor(0x90, 0xCD, 0xF4)  # Light Blue accent

    # Spacing
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(6)
    p_space.paragraph_format.space_after = Pt(6)

    # Metadata Summary Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    meta_table.columns[0].width = Inches(2.0)
    meta_table.columns[1].width = Inches(4.7)

    metadata = [
        ("Project:", "Customer Support AI Agent (LangGraph Workflow)"),
        ("Issue Diagnosed:", "User Profile Injected in Node but Resolving to 'None' in Final State"),
        ("Status:", "RESOLVED & FULLY VERIFIED"),
        ("Date / Release:", "October 2026 / Branch 1.1 Patch")
    ]

    for i, (label, val) in enumerate(metadata):
        c0 = meta_table.cell(i, 0)
        c1 = meta_table.cell(i, 1)
        set_cell_background(c0, "F2F4F7")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        
        # borders
        for cell in (c0, c1):
            set_cell_border(cell, 
                            bottom={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                            top={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                            left={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                            right={'val': 'single', 'sz': 4, 'color': 'E2E8F0'})

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.name = "Calibri"
        r0.font.size = Pt(10)
        r0.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(10)
        if label == "Status:":
            r1.bold = True
            r1.font.color.rgb = RGBColor(0x16, 0x65, 0x34)  # Forest Green
        else:
            r1.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(8)
    p_div.paragraph_format.space_after = Pt(8)

    # Helper for Section Headings
    def add_heading_1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(15)
        r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        return h

    def add_heading_2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
        return h

    def add_body(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = "Calibri"
            r_pre.font.size = Pt(10.5)
            r_pre.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = "Calibri"
            r_pre.font.size = Pt(10.5)
            r_pre.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        return p

    # =========================================================================
    # 1. EXECUTIVE SUMMARY
    # =========================================================================
    add_heading_1("1. Executive Summary")

    add_body(
        "In this project, a Customer Support Agent workflow is orchestrated using LangGraph to process user complaints "
        "through a sequence of specialized AI nodes: validation, context injection, categorization, sentiment analysis, "
        "priority assignment, email drafting, suggestion generation, priority reevaluation, and human escalation routing."
    )

    add_body(
        "During graph execution, the newly integrated node context_injection_node ran as expected and logged to the console: "
        "'🔧 SYSTEM: Injected profile for Alice (Tier: Enterprise)'. However, upon completion of the entire graph stream, "
        "the final summary output consistently displayed:",
        bold_prefix="The Symptom: "
    )

    add_code_block(doc, "🏁 Execution Complete. Checking Injected Memory:\n----------------------------------------\nUser Profile Found: None")

    add_body(
        "Why did this happen if the injection node logged success? In LangGraph, graph state is strictly managed through "
        "a centralized state schema—in this project, ComplaintState (defined as a Python TypedDict in Agent/state.py). "
        "When StateGraph(ComplaintState) compiles the execution DAG, it inspects ComplaintState.__annotations__ to create "
        "the graph's internal state channels. Only keys explicitly declared in the schema are preserved across state transitions.",
        bold_prefix="The Root Cause: "
    )

    add_body(
        "In the original codebase, ComplaintState only declared complaint_text, is_valid, validation_reason, category, "
        "sentiment, priority, priority_reason, response_email, suggestions, reevaluated_priority, escalation_required, and "
        "escalation_reason. Neither user_profile nor user_id was declared in ComplaintState."
    )

    add_callout_box(
        doc,
        "Because 'user_profile' and 'user_id' were absent from ComplaintState, LangGraph automatically filtered them out "
        "when writing updates back to the state channels after context_injection_node finished. Downstream nodes received "
        "a state dictionary stripped of the profile, and the final state dictionary had no record of it.",
        title="CORE DIAGNOSTIC FINDING"
    )

    add_body(
        "By updating ComplaintState in Agent/state.py to declare user_id and user_profile with appropriate type hints, "
        "LangGraph immediately recognized both channels. The user profile now persists through all graph nodes, successfully "
        "reaches human escalation and final output handlers, and prints the full profile dictionary.",
        bold_prefix="The Solution: "
    )

    # =========================================================================
    # 2. EXACT CHANGES MADE
    # =========================================================================
    add_heading_1("2. Exact Changes Made")

    add_body(
        "In adherence to the design constraint ('Do not rewrite my entire codebase or change my overall graph architecture; "
        "just fix the state-passing mechanism for the profile'), changes were isolated strictly to the state definition schema. "
        "No existing agent prompts, routing logic, or graph topologies were altered."
    )

    add_heading_2("File Modified: Agent/state.py")
    add_bullet("Imported Dict and Any from the standard typing library alongside TypedDict, Optional, and List.", bold_prefix="1. Type Imports: ")
    add_bullet("Added user_id: Optional[str] to the ComplaintState schema definition to store the originating customer ID.", bold_prefix="2. Added user_id: ")
    add_bullet("Added user_profile: Optional[Dict[str, Any]] to the ComplaintState schema definition to store customer metadata (name, tier, lifetime_value, recent_issue).", bold_prefix="3. Added user_profile: ")

    add_body(
        "Summary Table of File Modifications:",
        bold_prefix="Modification Inventory: "
    )

    # Table of modifications
    tbl_mods = doc.add_table(rows=2, cols=4)
    tbl_mods.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_mods.autofit = False
    tbl_mods.columns[0].width = Inches(1.8)
    tbl_mods.columns[1].width = Inches(1.2)
    tbl_mods.columns[2].width = Inches(1.0)
    tbl_mods.columns[3].width = Inches(2.7)

    headers = ["File Path", "Component", "Type of Change", "Description"]
    for col_idx, h in enumerate(headers):
        cell = tbl_mods.cell(0, col_idx)
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    row_data = [
        ("Customer-Support-Agent/Agent/state.py", "ComplaintState", "Schema Update", "Added user_id and user_profile channels with typing")
    ]
    for row_idx, data in enumerate(row_data, start=1):
        for col_idx, text in enumerate(data):
            cell = tbl_mods.cell(row_idx, col_idx)
            set_cell_background(cell, "F9FAFB")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            set_cell_border(cell, 
                            bottom={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                            top={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                            left={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                            right={'val': 'single', 'sz': 4, 'color': 'E2E8F0'})
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(6)
    p_sp2.paragraph_format.space_after = Pt(6)

    # =========================================================================
    # 3. CODE COMPARISON
    # =========================================================================
    add_heading_1("3. Code Comparison (Before vs. After)")

    add_body(
        "Below is the exact code comparison showing Agent/state.py before and after applying the fix."
    )

    add_heading_2("A. Original Code (Before Fix — Agent/state.py)")
    add_code_block(doc, 
"""# Original Agent/state.py
from typing import TypedDict, Optional, List

class ComplaintState(TypedDict):
    complaint_text: str

    is_valid: Optional[bool]
    validation_reason: Optional[str]

    category: Optional[str]
    sentiment: Optional[str]
    priority: Optional[str]
    priority_reason: Optional[str]

    response_email: Optional[str]
    suggestions: Optional[List[str]]

    reevaluated_priority: Optional[str]
    escalation_required: Optional[bool]
    escalation_reason: Optional[str]"""
    )

    add_heading_2("B. Fixed Code (After Fix — Agent/state.py)")
    add_code_block(doc, 
"""# Fixed Agent/state.py
from typing import TypedDict, Optional, List, Dict, Any

class ComplaintState(TypedDict):
    complaint_text: str
    user_id: Optional[str]
    user_profile: Optional[Dict[str, Any]]

    is_valid: Optional[bool]
    validation_reason: Optional[str]

    category: Optional[str]
    sentiment: Optional[str]
    priority: Optional[str]
    priority_reason: Optional[str]

    response_email: Optional[str]
    suggestions: Optional[List[str]]

    reevaluated_priority: Optional[str]
    escalation_required: Optional[bool]
    escalation_reason: Optional[str]"""
    )

    add_heading_2("C. Side-by-Side Unified Git Diff")
    add_code_block(doc,
"""--- a/Agent/state.py
+++ b/Agent/state.py
@@ -1,7 +1,9 @@
-from typing import TypedDict, Optional, List
+from typing import TypedDict, Optional, List, Dict, Any
 
 class ComplaintState(TypedDict):
     complaint_text: str
+    user_id: Optional[str]
+    user_profile: Optional[Dict[str, Any]]
 
     is_valid: Optional[bool]
     validation_reason: Optional[str]"""
    )

    add_body(
        "How this interacts with context_injection_node:",
        bold_prefix="Architectural Continuity: "
    )
    add_body(
        "In Agent/agents/user_validation.py, the context_injection_node reads state.get('user_id', 'user_123'), "
        "fetches the profile via get_user_profile(user_id), and sets state['user_profile'] = profile. "
        "Because user_profile is now an explicitly registered channel in ComplaintState, LangGraph's state manager "
        "allows the mutation to update the persistent state dictionary rather than dropping it at node exit."
    )

    # =========================================================================
    # 4. VERIFICATION STATUS
    # =========================================================================
    add_heading_1("4. Verification Status & Test Execution Results")

    add_body(
        "A complete end-to-end execution of test_graph.py was conducted using the project's dedicated Python virtual "
        "environment (Customer-Support-Agent/aiagent). The test executed the complete 10-node DAG.",
        bold_prefix="Test Execution: "
    )

    add_body(
        "Live Terminal Output Captured Post-Fix:",
        bold_prefix="Execution Log Verification: "
    )

    add_code_block(doc,
"""🚀 Starting Graph Execution...

DEBUG - Validation Output: is_valid=True reason='The user is reporting a payment and order issue (charged twice without order creation) and requesting a refund, which directly relates to e-commerce operations and expresses a support complaint.'
✅ Node Completed: [validate]
🔧 SYSTEM: Injected profile for Alice (Tier: Enterprise)
✅ Node Completed: [inject_context]
✅ Node Completed: [categorize]
✅ Node Completed: [sentiment]
✅ Node Completed: [priority]
✅ Node Completed: [response]
✅ Node Completed: [suggestion]
✅ Node Completed: [reevaluate]
✅ Node Completed: [escalation]

🚨 HUMAN ESCALATION REQUIRED 🚨

{'complaint_text': 'I was charged twice for order #99823, but the order was never created. I need a refund immediately!',
 'user_id': 'user_123',
 'user_profile': {'name': 'Alice', 'tier': 'Enterprise', 'lifetime_value': 15000, 'recent_issue': 'Server downtime on cluster B'},
 'is_valid': True,
 'validation_reason': 'The user is reporting a payment and order issue...',
 'category': 'payment_issue',
 'sentiment': 'angry',
 'priority': 'high',
 'priority_reason': 'Payment issue with angry sentiment',
 'response_email': 'Thank you for contacting us...',
 'suggestions': ['Escalate to billing team for urgent manual review', ...],
 'reevaluated_priority': 'high',
 'escalation_required': True,
 'escalation_reason': 'The customer was charged twice without an order being created...'}
✅ Node Completed: [human]

🏁 Execution Complete. Checking Injected Memory:
----------------------------------------
User Profile Found: {'name': 'Alice', 'tier': 'Enterprise', 'lifetime_value': 15000, 'recent_issue': 'Server downtime on cluster B'}
Category Assigned: payment_issue
Priority Assigned: high
Response Draft: Thank you for contacting us.

We have received your complaint and it is currently under review by our support team. An appropriate solution will be provided shortly.

We appreciate your patience.
----------------------------------------"""
    )

    add_callout_box(
        doc,
        "The summary output prints 'User Profile Found: {'name': 'Alice', 'tier': 'Enterprise', 'lifetime_value': 15000, "
        "'recent_issue': 'Server downtime on cluster B'}' instead of 'User Profile Found: None'. The bug is 100% resolved.",
        title="VERIFICATION CONFIRMED",
        border_color="166534",
        fill_color="F0FDF4"
    )

    # Save documents
    root_path_1 = os.path.abspath("Antigravity_Changelog.docx")
    root_path_2 = os.path.abspath(os.path.join("..", "Antigravity_Changelog.docx"))
    
    doc.save(root_path_1)
    print(f"Saved changelog document to: {root_path_1}")

    try:
        doc.save(root_path_2)
        print(f"Saved changelog document to: {root_path_2}")
    except Exception as e:
        print(f"Note on secondary path: {e}")

if __name__ == "__main__":
    build_changelog()

